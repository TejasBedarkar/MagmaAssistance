"""
LLM/document_reader.py

Turns ANY supported uploaded file into plain text the agent can reason over.

    images  : .jpg .jpeg .png .webp .gif .bmp .tif .tiff   -> GPT-4o Vision (OCR + description)
    PDF     : .pdf                                         -> native text layer, Vision OCR for scanned pages
    Excel   : .xlsx .xlsm (.xls needs `xlrd`)              -> every sheet as a table
    CSV/TSV : .csv .tsv                                    -> table (encoding + delimiter auto-detected)
    JSON    : .json .jsonl .ndjson                         -> pretty-printed + structure summary
    Word    : .docx                                        -> headings, paragraphs, tables (+ OCR of
                                                              embedded images when the doc is image-only)
    Text    : .txt .md .log

Detection is by FILE EXTENSION first and MIME type second, because browsers
(especially on Windows) report CSV as application/vnd.ms-excel or
application/octet-stream, which is exactly why a MIME-only allow-list rejects
perfectly good files.

Entry point: read_file(file_bytes, filename, mime_type, llm) -> dict
"""

import csv
import io
import json
import logging
import os
import zipfile
from datetime import date, datetime, time
from typing import Optional

logger = logging.getLogger("document-reader")

# --------------------------------------------------------------------------
# Limits (all overridable through env vars)
# --------------------------------------------------------------------------
MAX_STORED_CHARS = int(os.environ.get("DOC_MAX_STORED_CHARS", "500000"))
MAX_SHEET_ROWS = int(os.environ.get("DOC_MAX_SHEET_ROWS", "2000"))
MAX_SHEET_COLS = int(os.environ.get("DOC_MAX_SHEET_COLS", "60"))
MAX_CELL_CHARS = 500
MAX_JSON_ITEMS = 200
MAX_IMAGE_SIDE = 2048
MAX_DOCX_IMAGES_OCR = 5

# --------------------------------------------------------------------------
# File-type detection
# --------------------------------------------------------------------------
_EXT_TO_KIND = {
    ".pdf": "pdf",
    ".jpg": "image", ".jpeg": "image", ".png": "image", ".webp": "image",
    ".gif": "image", ".bmp": "image", ".tif": "image", ".tiff": "image",
    ".xlsx": "xlsx", ".xlsm": "xlsx",
    ".xls": "xls",
    ".csv": "csv", ".tsv": "csv",
    ".json": "json", ".jsonl": "json", ".ndjson": "json",
    ".docx": "docx",
    ".txt": "text", ".md": "text", ".log": "text",
}

_MIME_TO_KIND = {
    "application/pdf": "pdf",
    "image/jpeg": "image", "image/jpg": "image", "image/png": "image",
    "image/webp": "image", "image/gif": "image", "image/bmp": "image",
    "image/tiff": "image",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "application/vnd.ms-excel": "xls",
    "text/csv": "csv", "text/tab-separated-values": "csv",
    "application/json": "json", "application/x-ndjson": "json",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "text", "text/markdown": "text",
}

_VISION_NATIVE_MIME = {
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".webp": "image/webp", ".gif": "image/gif",
}

SUPPORTED_DESCRIPTION = (
    "images (JPG, PNG, WEBP, GIF, BMP, TIFF), PDF, Excel (XLSX/XLSM), CSV/TSV, "
    "JSON, Word (DOCX) and plain text (TXT/MD)"
)

# Accept string for <input type="file"> in the frontend
ACCEPT_ATTRIBUTE = (
    "image/*,.pdf,.csv,.tsv,.xlsx,.xlsm,.xls,.json,.jsonl,.docx,.txt,.md"
)


class DocumentReadError(Exception):
    """A file we recognise but could not read (corrupt, encrypted, too odd...)."""


def _ext(filename: Optional[str]) -> str:
    return os.path.splitext((filename or "").lower())[1]


def detect_kind(filename: Optional[str], mime_type: Optional[str] = None) -> Optional[str]:
    """Return 'pdf' | 'image' | 'xlsx' | 'xls' | 'csv' | 'json' | 'docx' | 'text',
    or None if the file type isn't supported. Extension wins over MIME."""
    kind = _EXT_TO_KIND.get(_ext(filename))
    if kind:
        return kind
    mime = (mime_type or "").split(";")[0].strip().lower()
    return _MIME_TO_KIND.get(mime)


# --------------------------------------------------------------------------
# Small formatting helpers
# --------------------------------------------------------------------------
def _fmt_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\r\n", " ").replace("\n", " ").replace("\r", " ").replace("\t", " ")
    text = text.replace("|", "/").strip()
    if len(text) > MAX_CELL_CHARS:
        text = text[:MAX_CELL_CHARS] + "…"
    return text


def _trim_row(row) -> list:
    cells = [_fmt_cell(v) for v in list(row)[:MAX_SHEET_COLS]]
    while cells and cells[-1] == "":
        cells.pop()
    return cells


def _rows_to_markdown(rows: list) -> str:
    """rows[0] is treated as the header row."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    norm = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(norm[0]) + " |", "|" + "|".join(["---"] * width) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in norm[1:]]
    return "\n".join(lines)


def _decode_bytes(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _result(text: str, kind: str, method: str, detail: str, page_count: int = 1,
            pages_read: Optional[int] = None) -> dict:
    truncated = False
    if len(text) > MAX_STORED_CHARS:
        text = text[:MAX_STORED_CHARS] + f"\n\n[... file truncated at {MAX_STORED_CHARS:,} characters ...]"
        truncated = True
    return {
        "text": text,
        "kind": kind,
        "method": method,
        "detail": detail,
        "page_count": page_count,
        "pages_read": page_count if pages_read is None else pages_read,
        "truncated": truncated,
    }


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------
def _prepare_image(file_bytes: bytes, filename: str) -> tuple:
    """Return (bytes, mime) ready for the Vision API.

    Fixes the things that silently break OCR on real-world uploads: EXIF
    rotation on phone photos, formats Vision doesn't accept (BMP/TIFF),
    transparency, and huge resolutions.
    """
    ext = _ext(filename)
    native_mime = _VISION_NATIVE_MIME.get(ext)
    try:
        from PIL import Image, ImageOps
    except ImportError:
        if native_mime:
            return file_bytes, native_mime
        raise DocumentReadError(
            f"Reading {ext or 'this'} images needs the Pillow package "
            "(pip install Pillow), or upload it as JPG/PNG instead."
        )

    try:
        img = Image.open(io.BytesIO(file_bytes))
        img.load()  # first frame only for animated GIF / multipage TIFF
        img = ImageOps.exif_transpose(img)
    except Exception as exc:
        raise DocumentReadError(f"Could not open this image (it may be corrupted): {exc}")

    if max(img.size) > MAX_IMAGE_SIDE:
        img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))

    out = io.BytesIO()
    if ext in (".jpg", ".jpeg"):
        img.convert("RGB").save(out, format="JPEG", quality=90)
        return out.getvalue(), "image/jpeg"

    if img.mode in ("RGBA", "LA", "P"):
        rgba = img.convert("RGBA")
        background = Image.new("RGB", rgba.size, (255, 255, 255))
        background.paste(rgba, mask=rgba.split()[-1])
        img = background
    else:
        img = img.convert("RGB")
    img.save(out, format="PNG")
    return out.getvalue(), "image/png"


def _read_image(file_bytes: bytes, filename: str, llm) -> dict:
    data, mime = _prepare_image(file_bytes, filename)
    extraction = llm.extract_document_text(file_bytes=data, mime_type=mime)
    text = (extraction.get("text") or "").strip()
    if not text:
        raise DocumentReadError("No readable content was found in this image.")
    return _result(text, "image", extraction.get("method", "vision"), "1 image")


def _read_pdf(file_bytes: bytes, llm) -> dict:
    extraction = llm.extract_document_text(file_bytes=file_bytes, mime_type="application/pdf")
    pages, read = extraction["page_count"], extraction["pages_read"]
    detail = f"{read}/{pages} page(s)" if read != pages else f"{pages} page(s)"
    res = _result(extraction["text"], "pdf", extraction["method"], detail, pages, read)
    if read < pages:
        res["truncated"] = True
    return res


def _read_xlsx(file_bytes: bytes) -> dict:
    try:
        from itertools import zip_longest
        from openpyxl import load_workbook
    except ImportError:
        raise DocumentReadError("Reading Excel files needs openpyxl (pip install openpyxl).")

    try:
        # Two streaming passes over the same sheets: cached VALUES, and the raw
        # cells (formulas). Files produced by scripts/other tools often have
        # formulas with no cached value -- fall back to showing the formula
        # text for those cells instead of silently dropping them.
        wb_values = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        wb_formulas = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=False)

        parts, total_rows, uncached = [], 0, 0
        for ws in wb_values.worksheets:
            ws_f = wb_formulas[ws.title]
            rows, seen = [], 0
            for raw_v, raw_f in zip_longest(ws.iter_rows(values_only=True), ws_f.iter_rows(values_only=True)):
                raw_v = raw_v or ()
                raw_f = raw_f or ()
                merged = []
                for v, f in zip_longest(raw_v, raw_f):
                    if v is None and isinstance(f, str) and f.startswith("="):
                        uncached += 1
                        v = f
                    merged.append(v)
                cells = _trim_row(merged)
                if not cells:
                    continue
                seen += 1
                if len(rows) < MAX_SHEET_ROWS:
                    rows.append(cells)
            total_rows += seen
            header = f"## Sheet: {ws.title}  ({seen} non-empty row(s))"
            if not rows:
                parts.append(f"{header}\n(empty)")
                continue
            body = _rows_to_markdown(rows)
            if seen > len(rows):
                body += f"\n[... {seen - len(rows)} more row(s) not shown ...]"
            parts.append(f"{header}\n{body}")
        n_sheets = len(wb_values.worksheets)
        wb_values.close()
        wb_formulas.close()
    except Exception as exc:
        raise DocumentReadError(
            f"Could not open this Excel file (it may be corrupted or password-protected): {exc}"
        )

    if not total_rows:
        raise DocumentReadError("This spreadsheet has no data in it.")
    text = "\n\n".join(parts)
    if uncached:
        text += (
            f"\n\n(Note: {uncached} formula cell(s) had no saved result in this file, "
            "so the formula itself is shown.)"
        )
    return _result(text, "xlsx", "openpyxl", f"{n_sheets} sheet(s), {total_rows} row(s)", n_sheets)


def _read_xls(file_bytes: bytes) -> dict:
    try:
        import pandas as pd
        sheets = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None, header=None, engine="xlrd")
    except ImportError:
        raise DocumentReadError(
            "Legacy .xls files need the 'xlrd' package. Please save the file as .xlsx and upload it again."
        )
    except Exception as exc:
        raise DocumentReadError(f"Could not open this .xls file: {exc}")

    parts, total_rows = [], 0
    for name, df in sheets.items():
        rows = []
        for raw in df.itertuples(index=False, name=None):
            cells = _trim_row([None if v != v else v for v in raw])  # NaN -> None
            if cells:
                rows.append(cells)
        total_rows += len(rows)
        shown = rows[:MAX_SHEET_ROWS]
        body = _rows_to_markdown(shown) if shown else "(empty)"
        if len(rows) > len(shown):
            body += f"\n[... {len(rows) - len(shown)} more row(s) not shown ...]"
        parts.append(f"## Sheet: {name}  ({len(rows)} non-empty row(s))\n{body}")
    if not total_rows:
        raise DocumentReadError("This spreadsheet has no data in it.")
    return _result("\n\n".join(parts), "xls", "xlrd", f"{len(sheets)} sheet(s), {total_rows} row(s)", len(sheets))


def _read_csv(file_bytes: bytes, filename: str) -> dict:
    text = _decode_bytes(file_bytes)
    if not text.strip():
        raise DocumentReadError("This file is empty.")
    csv.field_size_limit(10_000_000)
    if _ext(filename) == ".tsv":
        dialect = csv.excel_tab
    else:
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
    rows, seen = [], 0
    try:
        for raw in csv.reader(io.StringIO(text), dialect):
            cells = _trim_row(raw)
            if not cells:
                continue
            seen += 1
            if len(rows) < MAX_SHEET_ROWS:
                rows.append(cells)
    except csv.Error as exc:
        raise DocumentReadError(f"Could not parse this CSV: {exc}")
    if not rows:
        raise DocumentReadError("This file has no data rows.")
    body = _rows_to_markdown(rows)
    if seen > len(rows):
        body += f"\n[... {seen - len(rows)} more row(s) not shown ...]"
    out = f"## CSV data  ({seen} non-empty row(s), first row = header)\n{body}"
    return _result(out, "csv", "csv", f"{seen} row(s)")


def _describe_json(obj) -> str:
    if isinstance(obj, list):
        first = obj[0] if obj else None
        keys = f"; item keys: {', '.join(list(first.keys())[:25])}" if isinstance(first, dict) else ""
        return f"JSON array with {len(obj)} item(s){keys}"
    if isinstance(obj, dict):
        return f"JSON object with keys: {', '.join(list(obj.keys())[:40])}"
    return f"JSON value of type {type(obj).__name__}"


def _read_json(file_bytes: bytes, filename: str) -> dict:
    text = _decode_bytes(file_bytes).strip()
    if not text:
        raise DocumentReadError("This file is empty.")
    try:
        obj = json.loads(text)
    except json.JSONDecodeError:
        # JSON Lines: one object per line
        try:
            obj = [json.loads(line) for line in text.splitlines() if line.strip()]
        except json.JSONDecodeError as exc:
            raise DocumentReadError(f"This isn't valid JSON: {exc}")

    summary = _describe_json(obj)
    shown = obj
    if isinstance(obj, list) and len(obj) > MAX_JSON_ITEMS:
        shown = obj[:MAX_JSON_ITEMS]
        summary += f" (first {MAX_JSON_ITEMS} shown)"
    pretty = json.dumps(shown, indent=2, ensure_ascii=False, default=str)
    return _result(f"## {summary}\n{pretty}", "json", "json", summary)


def _read_text(file_bytes: bytes) -> dict:
    text = _decode_bytes(file_bytes).strip()
    if not text:
        raise DocumentReadError("This file is empty.")
    return _result(text, "text", "text", f"{len(text.splitlines())} line(s)")


def _read_docx(file_bytes: bytes, llm) -> dict:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        raise DocumentReadError("Reading Word files needs python-docx (pip install python-docx).")

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentReadError(
            f"Could not open this Word file (it may be corrupted, password-protected or an old .doc): {exc}"
        )

    blocks, n_tables = [], 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            if style.lower().startswith("heading"):
                digits = "".join(ch for ch in style if ch.isdigit())
                blocks.append("#" * min(int(digits or 1), 4) + " " + text)
            elif style.lower() == "title":
                blocks.append("# " + text)
            elif "list" in style.lower():
                blocks.append("- " + text)
            else:
                blocks.append(text)
        elif tag == "tbl":
            n_tables += 1
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells, prev = [], None
                for cell in row.cells:
                    if prev is not None and cell._tc is prev:  # merged cell repeats
                        continue
                    prev = cell._tc
                    cells.append(_fmt_cell(cell.text))
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append(_rows_to_markdown(rows))

    text = "\n\n".join(blocks).strip()
    method = "python-docx"
    detail = f"{len(blocks) - n_tables} paragraph(s), {n_tables} table(s)"

    # Image-only Word file (scanned pages pasted in): OCR the embedded pictures.
    if len(text) < 40 and llm is not None:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
            media = [n for n in zf.namelist()
                     if n.startswith("word/media/") and _ext(n) in _VISION_NATIVE_MIME]
            ocr_parts = []
            for i, name in enumerate(sorted(media)[:MAX_DOCX_IMAGES_OCR], 1):
                try:
                    data, mime = _prepare_image(zf.read(name), name)
                    ext = llm.extract_document_text(file_bytes=data, mime_type=mime)
                    if ext.get("text", "").strip():
                        ocr_parts.append(f"--- Embedded image {i} ---\n{ext['text'].strip()}")
                except Exception:
                    logger.exception("OCR of embedded docx image %s failed", name)
        if ocr_parts:
            text = (text + "\n\n" if text else "") + "\n\n".join(ocr_parts)
            method = "python-docx+vision"
            detail += f", {len(ocr_parts)} embedded image(s) OCR'd"

    if not text:
        raise DocumentReadError("No readable text was found in this Word file.")
    return _result(text, "docx", method, detail)


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------
def read_file(file_bytes: bytes, filename: str, mime_type: Optional[str], llm) -> dict:
    """Extract text from any supported file.

    `llm` is the LLM/VisionMixin instance (needed for images and scanned PDFs).

    Returns {"text", "kind", "method", "detail", "page_count", "pages_read",
    "truncated"}. Raises DocumentReadError for unsupported / unreadable files.
    """
    kind = detect_kind(filename, mime_type)
    if kind is None:
        raise DocumentReadError(
            f"Unsupported file type '{_ext(filename) or mime_type or 'unknown'}'. "
            f"Supported: {SUPPORTED_DESCRIPTION}."
        )

    logger.info("Reading '%s' as %s (%d bytes)", filename, kind, len(file_bytes))
    try:
        if kind == "image":
            return _read_image(file_bytes, filename, llm)
        if kind == "pdf":
            return _read_pdf(file_bytes, llm)
        if kind == "xlsx":
            return _read_xlsx(file_bytes)
        if kind == "xls":
            return _read_xls(file_bytes)
        if kind == "csv":
            return _read_csv(file_bytes, filename)
        if kind == "json":
            return _read_json(file_bytes, filename)
        if kind == "docx":
            return _read_docx(file_bytes, llm)
        return _read_text(file_bytes)
    except DocumentReadError:
        raise
    except RuntimeError as exc:  # raised by VisionMixin (corrupt / encrypted PDF etc.)
        raise DocumentReadError(str(exc))
